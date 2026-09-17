from collections import Counter, defaultdict
from datetime import datetime, time
from json import dumps
from pathlib import Path
from statistics import median
from unicodedata import normalize

from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "Asesinatos múltiples respaldo 2025 CH.xlsx"
OUTPUT = ROOT / "public" / "eventos.json"


def clean(value):
    if value is None:
        return None
    value = str(value).strip()
    return value if value and value.upper() not in {"NONE", "SIN_DATO"} else None


def mode(values):
    values = [clean(value) for value in values]
    values = [value for value in values if value]
    return Counter(values).most_common(1)[0][0] if values else None


def coordinate(value, minimum, maximum):
    try:
        number = float(value)
    except (TypeError, ValueError):
        return None
    for exponent in range(8):
        candidate = number / (10**exponent)
        if minimum <= candidate <= maximum:
            return round(candidate, 5)
    return None


def normalized_label(value):
    value = clean(value)
    if not value:
        return None
    replacements = {
        "LOS RIOS": "LOS RÍOS",
        "MANABI": "MANABÍ",
        "SANTO DOMINGO DE LOS TSACHILAS": "SANTO DOMINGO DE LOS TSÁCHILAS",
        "STO DGO DE LOS TSÁCHILAS": "SANTO DOMINGO DE LOS TSÁCHILAS",
    }
    return replacements.get(value, value)


workbook = load_workbook(SOURCE, read_only=True, data_only=True)
worksheet = workbook.active
groups = defaultdict(list)

for row in worksheet.iter_rows(min_row=2, values_only=True):
    if isinstance(row[19], datetime) and clean(row[0]):
        groups[clean(row[0])].append(row)

events = []
for code, rows in groups.items():
    dates = [row[19] for row in rows if isinstance(row[19], datetime)]
    latitudes = [coordinate(row[12], -6, 2) for row in rows]
    longitudes = [coordinate(row[14], -82, -74) for row in rows]
    latitudes = [value for value in latitudes if value is not None]
    longitudes = [value for value in longitudes if value is not None]
    ages = [row[27] for row in rows if isinstance(row[27], (int, float))]
    sexes = Counter(clean(row[29]) for row in rows)
    event_date = min(dates)
    event_time = next((row[21] for row in rows if isinstance(row[21], time)), None)
    events.append({
        "id": code,
        "date": event_date.date().isoformat(),
        "time": event_time.strftime("%H:%M") if event_time else None,
        "year": event_date.year,
        "victims": len(rows),
        "lat": round(median(latitudes), 5) if latitudes else None,
        "lng": round(median(longitudes), 5) if longitudes else None,
        "province": normalized_label(mode(row[8] for row in rows)),
        "canton": mode(row[10] for row in rows),
        "district": mode(row[4] for row in rows),
        "circuit": mode(row[5] for row in rows),
        "subcircuit": mode(row[7] for row in rows),
        "area": mode(row[16] for row in rows),
        "place": mode(row[17] for row in rows),
        "placeType": mode(row[18] for row in rows),
        "weapon": mode(row[22] for row in rows),
        "weaponType": mode(row[23] for row in rows),
        "motivation": mode(row[24] for row in rows),
        "observedMotivation": mode(row[25] for row in rows),
        "ageMin": int(min(ages)) if ages else None,
        "ageMax": int(max(ages)) if ages else None,
        "men": sexes.get("HOMBRE", 0),
        "women": sexes.get("MUJER", 0),
    })

events.sort(key=lambda event: (event["date"], event["id"]), reverse=True)
OUTPUT.write_text(dumps(events, ensure_ascii=False, separators=(",", ":")), encoding="utf-8")
print(f"{len(events)} eventos; {sum(e['victims'] for e in events)} vítimas; {sum(e['lat'] is not None for e in events)} georreferenciados")
